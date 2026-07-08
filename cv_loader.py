"""
cv_loader.py — Chargement du CV depuis un fichier PDF ou DOCX
"""

import io
from typing import Optional


def load_cv_from_pdf(file_bytes: bytes) -> str:
    """Extrait le texte d'un fichier PDF."""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n\n".join(text_parts)
    except ImportError:
        raise ImportError("pdfplumber n'est pas installé. Lancez : pip install pdfplumber")
    except Exception as e:
        raise ValueError(f"Impossible de lire le PDF : {e}")


def load_cv_from_docx(file_bytes: bytes) -> str:
    """Extrait le texte d'un fichier DOCX."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        # Extraire également les tableaux (souvent utilisés dans les CV)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        return "\n".join(paragraphs)
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")
    except Exception as e:
        raise ValueError(f"Impossible de lire le DOCX : {e}")


def load_cv_from_txt(file_bytes: bytes) -> str:
    """Extrait le texte d'un fichier TXT."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1")


def load_cv(file_bytes: bytes, filename: str) -> str:
    """
    Charge le contenu d'un CV à partir de ses bytes et de son nom de fichier.
    Supporte : PDF, DOCX, DOC, TXT
    """
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        return load_cv_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        return load_cv_from_docx(file_bytes)
    elif filename_lower.endswith(".doc"):
        # Tentative de lecture comme DOCX, sinon comme texte brut
        try:
            return load_cv_from_docx(file_bytes)
        except Exception:
            return load_cv_from_txt(file_bytes)
    elif filename_lower.endswith(".txt"):
        return load_cv_from_txt(file_bytes)
    else:
        raise ValueError(f"Format de fichier non supporté : {filename}. Utilisez PDF, DOCX, DOC ou TXT.")


def generate_docx_from_text(title: str, content: str) -> bytes:
    """
    Génère un fichier DOCX à partir d'un titre et d'un texte.
    Retourne les bytes du fichier DOCX.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Style de base
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Titre
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x2E, 0x4A, 0x7A)

    doc.add_paragraph()  # Espace

    # Contenu : on split sur les lignes vides pour créer des paragraphes
    sections = content.split("\n\n")
    for section in sections:
        lines = section.strip().split("\n")
        if not lines:
            continue

        first_line = lines[0].strip()

        # Détection des titres de section (ligne courte en majuscules ou avec ###)
        if (first_line.startswith("###") or
                first_line.startswith("##") or
                (len(first_line) < 60 and first_line.isupper()) or
                first_line.endswith(":")):
            clean_title = first_line.lstrip("#").strip().rstrip(":")
            h = doc.add_heading(clean_title, level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x2E, 0x4A, 0x7A)
            # Reste de la section
            rest = "\n".join(lines[1:]).strip()
            if rest:
                para = doc.add_paragraph(rest)
                para.paragraph_format.space_after = Pt(6)
        else:
            full_text = "\n".join(lines).strip()
            para = doc.add_paragraph(full_text)
            para.paragraph_format.space_after = Pt(6)

    # Sauvegarde en mémoire
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
