"""
AutoCandidature IA — Version tout-en-un
Toute la logique dans un seul fichier, aucune dépendance interne.
"""

import io
import re
import streamlit as st
import anthropic

# ─── Chargement du CV ─────────────────────────────────────────────────────────

def extract_pdf(data: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n\n".join(p.extract_text() or "" for p in pdf.pages)
    except ImportError:
        pass
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        return ""

def extract_docx(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return ""

def load_cv(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith(".pdf"):
        return extract_pdf(data)
    elif name.endswith(".docx"):
        return extract_docx(data)
    else:
        return data.decode("utf-8", errors="ignore")

def export_docx(content: str) -> bytes | None:
    try:
        import docx as dx
        from docx.shared import Pt
        doc = dx.Document()
        for line in content.split("\n"):
            s = line.strip()
            if s.startswith("# "):
                doc.add_heading(s[2:], level=1)
            elif s.startswith("## "):
                doc.add_heading(s[3:], level=2)
            else:
                p = doc.add_paragraph()
                parts = re.split(r"(\*\*[^*]+\*\*)", s)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()
    except Exception:
        return None

# ─── Adaptation IA ────────────────────────────────────────────────────────────

def stream_adapted_cv(cv: str, job: str, sector: str, api_key: str, model: str):
    client = anthropic.Anthropic(api_key=api_key)
    with client.messages.stream(
        model=model,
        max_tokens=4096,
        system="""Tu es un expert en recrutement et rédaction de CV professionnels.
Adapte le CV fourni pour le poste cible :
- Conserve toutes les informations réelles (ne rien inventer)
- Réorganise et reformule pour coller aux exigences du poste
- Utilise les mots-clés de l'offre
- Structure avec des sections claires : Profil, Expériences, Compétences, Formation
- Réponds en français
- Utilise # pour le nom/titre, ## pour les sections, ** pour les éléments clés""",
        messages=[{"role": "user", "content": f"""Mon CV :\n\n{cv}\n\n---\n\nOffre d'emploi (secteur : {sector}) :\n\n{job}\n\n---\n\nAdapte mon CV pour ce poste. Garde mes vraies informations, mets en avant ce qui est pertinent. Retourne uniquement le CV adapté."""}],
    ) as stream:
        for text in stream.text_stream:
            yield text

def stream_cover_letter(cv: str, job: str, sector: str, company: str, name: str, api_key: str, model: str):
    client = anthropic.Anthropic(api_key=api_key)
    company_info = f" chez {company}" if company else ""
    name_info = f"Le candidat s'appelle {name}." if name else "Utilise [Votre Prénom Nom] pour le nom."
    with client.messages.stream(
        model=model,
        max_tokens=2048,
        system="""Tu es un expert en rédaction de lettres de motivation percutantes et professionnelles.
La lettre doit :
- Avoir une accroche originale liée au poste
- Montrer 2-3 réalisations clés en lien avec l'offre
- Avoir un ton chaleureux mais professionnel
- Faire 3-4 paragraphes (300-400 mots)
- Se terminer par un appel à l'entretien
- Inclure la date, l'objet et une formule de politesse
- Être rédigée en français""",
        messages=[{"role": "user", "content": f"""CV :\n\n{cv}\n\n---\n\nOffre{company_info} (secteur : {sector}) :\n\n{job}\n\n---\n\n{name_info}\n\nRédige une lettre de motivation personnalisée et percutante."""}],
    ) as stream:
        for text in stream.text_stream:
            yield text

# ─── Interface Streamlit ──────────────────────────────────────────────────────

st.set_page_config(
    page_title="AutoCandidature IA",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.main-title { font-size: 2.2rem; font-weight: 700; color: #1a1a2e; margin-bottom: 0.2rem; }
.subtitle { font-size: 1rem; color: #6c757d; margin-bottom: 2rem; }
.section-header { font-size: 1.1rem; font-weight: 600; color: #16213e;
  border-bottom: 2px solid #e9ecef; padding-bottom: 0.3rem; margin-bottom: 1rem; }
</style>
""", unsafe_allow_html=True)

st.markdown('<p class="main-title">📄 AutoCandidature IA</p>', unsafe_allow_html=True)
st.markdown('<p class="subtitle">Adaptez votre CV et générez une lettre de motivation personnalisée pour chaque offre d\'emploi.</p>', unsafe_allow_html=True)

# ── Barre latérale ─────────────────────────────────────────────────────────────

with st.sidebar:
    st.header("⚙️ Configuration")

    # Clé API : d'abord depuis les secrets Streamlit, sinon champ manuel
    default_key = ""
    try:
        default_key = st.secrets.get("ANTHROPIC_API_KEY", "")
    except Exception:
        pass

    if default_key:
        api_key = default_key
        st.success("🔑 Clé API chargée depuis les secrets")
    else:
        api_key = st.text_input(
            "Clé API Anthropic",
            type="password",
            placeholder="sk-ant-...",
            help="Obtenez votre clé sur console.anthropic.com",
        )

    model = st.selectbox(
        "Modèle Claude",
        options=["claude-sonnet-4-6", "claude-opus-4-6", "claude-haiku-4-5-20251001"],
        index=0,
        help="Sonnet = équilibre qualité/rapidité | Opus = meilleure qualité | Haiku = plus rapide",
    )

    st.divider()
    st.header("👤 Informations candidat")
    applicant_name = st.text_input("Votre nom complet", placeholder="Jean Dupont")
    company_name = st.text_input("Nom de l'entreprise", placeholder="Google, LVMH...")

    st.divider()
    st.markdown("### ℹ️ À propos")
    st.markdown("Cet outil utilise **Claude (Anthropic)** pour adapter votre CV et générer une lettre personnalisée.\n\n**Vos données ne sont pas stockées.**")

# ── Zone principale ─────────────────────────────────────────────────────────────

col1, col2 = st.columns([1, 1], gap="large")

with col1:
    st.markdown('<p class="section-header">1. Votre CV</p>', unsafe_allow_html=True)
    cv_source = st.radio("Source du CV", ["Uploader un fichier", "Coller le texte"], horizontal=True, label_visibility="collapsed")
    cv_content = ""

    if cv_source == "Uploader un fichier":
        uploaded = st.file_uploader("Choisissez votre CV", type=["pdf", "docx", "txt"])
        if uploaded:
            try:
                cv_content = load_cv(uploaded)
                st.success(f"✅ CV chargé ({len(cv_content)} caractères)")
                with st.expander("Aperçu"):
                    st.text(cv_content[:1500] + ("..." if len(cv_content) > 1500 else ""))
            except Exception as e:
                st.error(f"Erreur : {e}")
    else:
        cv_content = st.text_area("Collez votre CV ici", height=300, placeholder="Nom Prénom\nTitre professionnel\n\nEXPÉRIENCES\n...")

with col2:
    st.markdown('<p class="section-header">2. Offre d\'emploi</p>', unsafe_allow_html=True)
    sector = st.text_input("Secteur / Domaine", placeholder="Ex: Data Science, Marketing, Finance...")
    job_posting = st.text_area("Collez l'offre d'emploi complète", height=300,
        placeholder="Titre : ...\nEntreprise : ...\nDescription : ...\nProfil recherché : ...")

# ── Bouton ─────────────────────────────────────────────────────────────────────

st.divider()
can_generate = bool(api_key and cv_content and job_posting and sector)
generate_btn = st.button("🚀 Générer le CV adapté et la lettre de motivation",
    type="primary", use_container_width=True, disabled=not can_generate)

if not api_key:
    st.caption("⚠️ Entrez votre clé API Anthropic dans la barre latérale.")
elif not cv_content:
    st.caption("⚠️ Chargez ou collez votre CV.")
elif not sector:
    st.caption("⚠️ Précisez le secteur / domaine.")
elif not job_posting:
    st.caption("⚠️ Collez l'offre d'emploi.")

# ── Génération ─────────────────────────────────────────────────────────────────

if generate_btn:
    col_cv, col_letter = st.columns(2, gap="large")
    full_cv = ""
    full_letter = ""

    with col_cv:
        st.markdown("### 📋 CV Adapté")
        cv_box = st.empty()
        try:
            for chunk in stream_adapted_cv(cv_content, job_posting, sector, api_key, model):
                full_cv += chunk
                cv_box.markdown(full_cv)
            st.session_state["adapted_cv"] = full_cv
            st.success("✅ CV adapté généré !")
        except Exception as e:
            st.error(f"Erreur : {e}")

    with col_letter:
        st.markdown("### ✉️ Lettre de Motivation")
        letter_box = st.empty()
        try:
            for chunk in stream_cover_letter(cv_content, job_posting, sector, company_name, applicant_name, api_key, model):
                full_letter += chunk
                letter_box.markdown(full_letter)
            st.session_state["cover_letter"] = full_letter
            st.success("✅ Lettre générée !")
        except Exception as e:
            st.error(f"Erreur : {e}")

    # ── Téléchargements ────────────────────────────────────────────────────────
    adapted_cv = st.session_state.get("adapted_cv", "")
    cover_letter = st.session_state.get("cover_letter", "")

    if adapted_cv and cover_letter:
        st.divider()
        st.markdown("### 💾 Télécharger les documents")
        c1, c2, c3, c4 = st.columns(4)

        cv_docx = export_docx(adapted_cv)
        letter_docx = export_docx(cover_letter)

        with c1:
            if cv_docx:
                st.download_button("📥 CV (.docx)", data=cv_docx, file_name="cv_adapte.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True)
        with c2:
            st.download_button("📥 CV (.txt)", data=adapted_cv.encode(), file_name="cv_adapte.txt",
                mime="text/plain", use_container_width=True)
        with c3:
            if letter_docx:
                st.download_button("📥 Lettre (.docx)", data=letter_docx, file_name="lettre_motivation.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True)
        with c4:
            st.download_button("📥 Lettre (.txt)", data=cover_letter.encode(), file_name="lettre_motivation.txt",
                mime="text/plain", use_container_width=True)

elif st.session_state.get("adapted_cv") and st.session_state.get("cover_letter"):
    col_cv, col_letter = st.columns(2, gap="large")
    with col_cv:
        st.markdown("### 📋 CV Adapté")
        st.markdown(st.session_state["adapted_cv"])
    with col_letter:
        st.markdown("### ✉️ Lettre de Motivation")
        st.markdown(st.session_state["cover_letter"])
